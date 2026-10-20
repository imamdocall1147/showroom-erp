# -*- coding: utf-8 -*-
"""
data_cleaned.xlsx থেকে পুরনো ক্রয় (স্টক) ও বিক্রয়ের রেকর্ড সিস্টেমে ঢোকায়।
• সব গাড়ী একটা মডেল ও একটা গোডাউনে ধরা হয়, দাম ০ থাকে, কোনো হিসাবের (ভাউচার) এন্ট্রি হয় না।
• প্রথমে "ট্রায়াল" চলে (কিছু লেখে না), ঠিক লাগলে --commit দিয়ে আসলটা চালান।

ব্যবহার (PowerShell, backend ফোল্ডারে, venv চালু করে):
    python import_data.py ..\\data_cleaned.xlsx                      # মডেল ও গোডাউনের তালিকা দেখায়
    python import_data.py ..\\data_cleaned.xlsx --product-id 3 --warehouse-id 2           # ট্রায়াল
    python import_data.py ..\\data_cleaned.xlsx --product-id 3 --warehouse-id 2 --commit  # আসল ঢোকানো
"""
import argparse
import collections
import datetime
import os
import sys

import openpyxl


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def as_dt(d, fallback):
    """তারিখ -> ডাটাবেসের সময় (UTC ০৬:০০ = বাংলাদেশে দুপুর ১২টা, তাই তারিখ সরে যায় না)"""
    d = d or fallback
    return datetime.datetime.combine(d, datetime.time(6, 0))


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    purchases, sales = [], []
    ws = wb["ক্রয়"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        purchases.append({
            "row": r[0],
            "date": as_date(r[2]),
            "engine": str(r[3] or "").strip(),
            "chassis": str(r[4] or "").strip(),
        })
    ws = wb["বিক্রয়"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        sales.append({
            "row": r[0],
            "chassis": str(r[1] or "").strip(),
            "engine": str(r[2] or "").strip(),
            "date": as_date(r[3]),
            "customer": str(r[5] or "").strip(),
            "address": str(r[6] or "").strip(),
            "ref": str(r[8] or "").strip(),
        })
    return purchases, sales


def valid_chassis(c):
    return bool(c) and c != "#N/A" and len(c) == 17


def plan(purchases, sales, existing_chassis, existing_engine, create_missing=True):
    """কী কী ঢুকবে তার পরিকল্পনা (ডাটাবেসে হাত না দিয়ে)"""
    notes = collections.defaultdict(list)
    units = collections.OrderedDict()      # chassis -> dict
    engine_to_chassis = {}
    for e, c in existing_engine.items():
        engine_to_chassis[e] = c

    for p in purchases:
        c, e = p["chassis"], p["engine"]
        if not valid_chassis(c) and len(c) != 18:
            notes["ক্রয়: চেসিস ঠিক নেই, বাদ"].append(p["row"]); continue
        if c in existing_chassis:
            notes["ক্রয়: আগে থেকেই সিস্টেমে আছে, বাদ"].append(p["row"]); continue
        if c in units:
            notes["ক্রয়: চেসিস ফাইলে দুইবার, পরেরটা বাদ"].append(p["row"]); continue
        if e in existing_engine or e in engine_to_chassis:
            notes["ক্রয়: ইঞ্জিন নম্বর আগে আছে, বাদ"].append(p["row"]); continue
        units[c] = {"chassis": c, "engine": e, "date": p["date"], "sold": None, "from": "purchase"}
        engine_to_chassis[e] = c

    sold_count = 0
    for s in sales:
        c, e = s["chassis"], s["engine"]
        if not valid_chassis(c):
            c = engine_to_chassis.get(e, "") if e else ""
            if not c:
                notes["বিক্রয়: চেসিস ঠিক নেই ও ইঞ্জিন দিয়েও পাওয়া যায়নি, বাদ"].append(s["row"]); continue
        if c in existing_chassis and c not in units:
            notes["বিক্রয়: গাড়ী আগে থেকেই সিস্টেমে আছে, বাদ"].append(s["row"]); continue
        if c not in units:
            if create_missing and e and e not in engine_to_chassis:
                units[c] = {"chassis": c, "engine": e, "date": s["date"], "sold": None, "from": "sale-only"}
                engine_to_chassis[e] = c
                notes["বিক্রয়: ক্রয়ের তালিকায় ছিল না, গাড়ী নিজে বানানো হলো"].append(s["row"])
            else:
                notes["বিক্রয়: গাড়ী খুঁজে পাওয়া যায়নি, বাদ"].append(s["row"]); continue
        if units[c]["sold"] is not None:
            notes["বিক্রয়: একই গাড়ী আবার বিক্রি, পরেরটা বাদ"].append(s["row"]); continue
        units[c]["sold"] = s
        sold_count += 1

    customers = collections.OrderedDict()
    for u in units.values():
        s = u["sold"]
        if not s or not s["customer"]:
            continue
        e = customers.setdefault(s["customer"], {"addr": collections.Counter(), "ref": collections.Counter(), "first": None, "n": 0})
        e["n"] += 1
        if s["address"]:
            e["addr"][s["address"]] += 1
        if s["ref"]:
            e["ref"][s["ref"]] += 1
        if s["date"] and (e["first"] is None or s["date"] < e["first"]):
            e["first"] = s["date"]
    return units, customers, notes, sold_count


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--product-id", type=int)
    ap.add_argument("--warehouse-id", type=int)
    ap.add_argument("--commit", action="store_true", help="আসলে ডাটাবেসে লিখুন (না দিলে শুধু ট্রায়াল)")
    ap.add_argument("--no-create-missing", action="store_true", help="ক্রয়ের তালিকায় নেই এমন বিক্রি গাড়ী বাদ দিন")
    ap.add_argument("--selftest", action="store_true", help="ডাটাবেস ছাড়া শুধু ফাইল পরীক্ষা")
    a = ap.parse_args()

    purchases, sales = read_workbook(a.xlsx)
    print(f"ফাইলে ক্রয় {len(purchases)} সারি, বিক্রয় {len(sales)} সারি")

    if a.selftest:
        units, customers, notes, sold = plan(purchases, sales, {}, {}, not a.no_create_missing)
        print(f"[ডাটাবেস ছাড়া পরীক্ষা] গাড়ী {len(units)}, বিক্রি {sold}, স্টকে {len(units) - sold}, কাস্টমার {len(customers)}")
        for k, v in notes.items():
            print(f"  - {k}: {len(v)}")
        return

    sys.path.insert(0, os.getcwd())
    from sqlalchemy.orm import Session
    from database import engine
    from models import Product, Warehouse, Party, Unit

    with Session(engine) as db:
        if not a.product_id or not a.warehouse_id:
            print("\nমডেলের তালিকা (product-id):")
            for p in db.query(Product).order_by(Product.id).all():
                print(f"  {p.id}: {p.model_name}")
            print("\nগোডাউনের তালিকা (warehouse-id):")
            for w in db.query(Warehouse).order_by(Warehouse.id).all():
                print(f"  {w.id}: {w.name} (শোরুম নং {w.branch_id})")
            print("\nউপরের তালিকা থেকে নম্বর বেছে --product-id ও --warehouse-id দিয়ে আবার চালান।")
            return

        product = db.query(Product).filter(Product.id == a.product_id).first()
        warehouse = db.query(Warehouse).filter(Warehouse.id == a.warehouse_id).first()
        if not product or not warehouse:
            print("[ত্রুটি] মডেল বা গোডাউনের নম্বর ভুল।")
            return

        existing_chassis = {c for (c,) in db.query(Unit.chassis_no).all()}
        existing_engine = {e: c for (e, c) in db.query(Unit.engine_no, Unit.chassis_no).all()}
        units, customers, notes, sold = plan(purchases, sales, existing_chassis, existing_engine, not a.no_create_missing)

        print(f"\nমডেল: {product.model_name} | গোডাউন: {warehouse.name}")
        print(f"নতুন গাড়ী ঢুকবে: {len(units)} (এর মধ্যে বিক্রি {sold}, স্টকে {len(units) - sold})")
        print(f"কাস্টমার (ফাইল অনুযায়ী): {len(customers)}")
        for k, v in notes.items():
            print(f"  - {k}: {len(v)} (সারি: {v[:8]}{' ...' if len(v) > 8 else ''})")

        if not a.commit:
            print("\nএটা শুধু ট্রায়াল ছিল, কিছু লেখা হয়নি। ঠিক লাগলে শেষে --commit যোগ করে আবার চালান।")
            return

        has_ref = hasattr(Party, "reference")
        has_entry = hasattr(Party, "entry_date")
        has_sale_ref = hasattr(Unit, "sale_ref")
        if not (has_ref and has_entry and has_sale_ref):
            print("[সতর্কতা] নতুন কলামগুলো (reference, entry_date, sale_ref) models.py-তে নেই। "
                  "আগে apply_partners.py চালান, নইলে রেফারেন্স ঢুকবে না। এখনকার মতো এগিয়ে যাচ্ছি না।")
            return

        existing_parties = {p.name: p for p in db.query(Party).filter(Party.party_type == "customer").all()}
        now = datetime.date.today()
        created_parties = 0
        for name, e in customers.items():
            party = existing_parties.get(name)
            if not party:
                party = Party(
                    name=name,
                    party_type="customer",
                    address=(e["addr"].most_common(1)[0][0] if e["addr"] else None),
                    reference=(e["ref"].most_common(1)[0][0] if e["ref"] else None),
                    entry_date=e["first"],
                )
                db.add(party)
                existing_parties[name] = party
                created_parties += 1
        db.flush()

        for u in units.values():
            s = u["sold"]
            unit = Unit(
                product_id=product.id,
                warehouse_id=warehouse.id,
                chassis_no=u["chassis"],
                engine_no=u["engine"],
                purchase_price=0,
                purchase_date=as_dt(u["date"], now),
                status="in_stock",
            )
            if s:
                unit.status = "sold"
                unit.sold_price = 0
                unit.decoration_cost = 0
                unit.registration_cost = 0
                unit.sold_date = as_dt(s["date"], now)
                if s["customer"]:
                    unit.customer_id = existing_parties[s["customer"]].id
                if s["ref"]:
                    unit.sale_ref = s["ref"]
            db.add(unit)
        db.commit()
        print(f"\nসফল! {len(units)} টি গাড়ী ও {created_parties} জন নতুন কাস্টমার ঢোকানো হয়েছে।")


if __name__ == "__main__":
    main()
